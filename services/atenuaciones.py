import io

import pandas as pd
from openpyxl.styles import PatternFill


def colorear_prioridad(ws):
    colores = {
        "ALTA": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "MEDIA": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "BAJA": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    }

    prioridad_col = None
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == "prioridad":
            prioridad_col = idx
            break

    if not prioridad_col:
        return

    for row in ws.iter_rows(min_row=2):
        valor = row[prioridad_col - 1].value
        if valor in colores:
            for cell in row:
                cell.fill = colores[valor]


def determinar_danio_dominante(row):
    los = row.get("onus_los", 0)
    power = row.get("onus_power_fail", 0)

    if los > power and los > 0:
        return "FIBRA"
    if power > los and power > 0:
        return "ENERGIA"
    if los == 0 and power == 0:
        return "SIN_CORTE"
    return "MIXTO"


PESOS_ESTADO = {
    "Arpón Crítico": 3,
    "Arpón Alarmado": 2,
}


def clasificar_prioridad(row):
    if bool(row.get("puerto_muerto", False)):
        return "ALTA"

    score = row.get("score_prioridad", 0)
    if score >= 100:
        return "ALTA"
    if score >= 50:
        return "MEDIA"
    return "BAJA"


def clasificar_estado_1310(valor):
    if pd.isna(valor):
        return "Normal"
    if valor <= -33:
        return "Arpón Crítico"
    if valor <= -30:
        return "Arpón Alarmado"
    return "Normal"


def clasificar_estado_1490(valor):
    if pd.isna(valor):
        return "Normal"
    if valor <= -30:
        return "Arpón Crítico"
    if valor <= -28:
        return "Arpón Alarmado"
    return "Normal"


def analizar_senal(df, olt, columna_signal):
    df_olt = df[df["OLT"] == olt].copy()

    resumen = (
        df_olt.groupby(["OLT", "Board", "Port"])
        .agg(
            promedio_1310=("Signal 1310", "mean"),
            promedio_1490=("Signal 1490", "mean"),
            total_onus=("ONU external ID", "count"),
            onus_online=("Status", lambda x: (x == "Online").sum()),
            onus_offline=("Status", lambda x: (x == "Offline").sum()),
            onus_los=("Status", lambda x: (x == "LOS").sum()),
            onus_power_fail=("Status", lambda x: (x == "Power fail").sum()),
            onus_criticas=(columna_signal, lambda x: (x <= -30).sum()),
            address=("Address", lambda x: x.mode().iloc[0] if not x.mode().empty else None),
        )
        .reset_index()
    )

    if columna_signal == "Signal 1310":
        resumen["senal_evaluada"] = "1310"
        resumen["estado"] = resumen["promedio_1310"].apply(clasificar_estado_1310)
    else:
        resumen["senal_evaluada"] = "1490"
        resumen["estado"] = resumen["promedio_1490"].apply(clasificar_estado_1490)

    resumen = resumen[resumen["estado"] != "Normal"].copy()
    if resumen.empty:
        return resumen

    resumen["peso_estado"] = resumen["estado"].map(PESOS_ESTADO).fillna(0)
    resumen["score_prioridad"] = resumen["total_onus"] * resumen["peso_estado"]
    resumen["porcentaje_online"] = ((resumen["onus_online"] / resumen["total_onus"]) * 100).round(1)
    resumen["puerto_muerto"] = resumen["onus_online"] == 0
    resumen = resumen.sort_values(by="score_prioridad", ascending=False)
    return resumen


def generar_agenda_general(df):
    agendas = []

    for olt in sorted(df["OLT"].dropna().unique()):
        for signal in ["Signal 1310", "Signal 1490"]:
            agenda_olt = analizar_senal(df, olt, signal)
            if not agenda_olt.empty:
                agendas.append(agenda_olt)

    if not agendas:
        return pd.DataFrame()

    agenda_general = pd.concat(agendas, ignore_index=True)
    agenda_general["prioridad"] = agenda_general.apply(clasificar_prioridad, axis=1)
    agenda_general["danio_dominante"] = agenda_general.apply(determinar_danio_dominante, axis=1)

    columnas = [
        "prioridad", "danio_dominante", "score_prioridad",
        "OLT", "Board", "Port", "senal_evaluada",
        "estado", "promedio_1310", "promedio_1490", "total_onus",
        "onus_online", "onus_offline", "onus_los", "onus_power_fail",
        "porcentaje_online", "puerto_muerto", "address",
    ]
    agenda_general = agenda_general[columnas].copy()

    orden_prioridad = {"ALTA": 0, "MEDIA": 1, "BAJA": 2}
    agenda_general["orden_prioridad"] = agenda_general["prioridad"].map(orden_prioridad)
    agenda_general = agenda_general.sort_values(
        by=["orden_prioridad", "score_prioridad"],
        ascending=[True, False],
    ).drop(columns=["orden_prioridad"])
    return agenda_general


def separar_agendas_por_danio(agenda_general):
    if agenda_general.empty:
        return {}

    return {
        "AGENDA_FIBRA": agenda_general[agenda_general["danio_dominante"] == "FIBRA"],
        "AGENDA_ENERGIA": agenda_general[agenda_general["danio_dominante"] == "ENERGIA"],
        "AGENDA_MIXTA": agenda_general[agenda_general["danio_dominante"] == "MIXTO"],
        "AGENDA_SIN_CORTE": agenda_general[agenda_general["danio_dominante"] == "SIN_CORTE"],
    }


def generar_excel_atenuaciones(df):
    output = io.BytesIO()
    hojas_prioridad = []
    hojas_creadas = False

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for olt in sorted(df["OLT"].dropna().unique()):
            for signal in ["Signal 1310", "Signal 1490"]:
                resumen = analizar_senal(df, olt, signal)
                if resumen.empty:
                    continue

                hoja = f"{olt}_{signal.split()[-1]}"[:31]
                resumen.to_excel(writer, sheet_name=hoja, index=False)
                hojas_creadas = True

        agenda = generar_agenda_general(df)
        if not agenda.empty:
            agenda.to_excel(writer, sheet_name="AGENDA_GENERAL", index=False)
            hojas_prioridad.append("AGENDA_GENERAL")
            hojas_creadas = True

            agendas_sep = separar_agendas_por_danio(agenda)
            for nombre, df_ag in agendas_sep.items():
                if not df_ag.empty:
                    df_ag.to_excel(writer, sheet_name=nombre, index=False)
                    hojas_prioridad.append(nombre)

        if not hojas_creadas:
            pd.DataFrame(
                {
                    "mensaje": [
                        "No se encontraron arpónes críticos ni alarmados",
                        "Todas las señales están en estado NORMAL",
                    ]
                }
            ).to_excel(writer, sheet_name="SIN_ALERTAS", index=False)

        workbook = writer.book
        for nombre_hoja in hojas_prioridad:
            if nombre_hoja in workbook.sheetnames:
                colorear_prioridad(workbook[nombre_hoja])

    output.seek(0)
    return output


def procesar_atenuaciones(csv_file):
    df = pd.read_csv(csv_file)
    columnas_requeridas = [
        "OLT", "Board", "Port", "Signal 1310", "Signal 1490",
        "ONU external ID", "Status", "Address",
    ]
    faltantes = [col for col in columnas_requeridas if col not in df.columns]
    if faltantes:
        raise ValueError(
            f"El archivo CSV no contiene las columnas requeridas: {', '.join(faltantes)}"
        )

    for col in ["Signal 1310", "Signal 1490"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    agenda_general = generar_agenda_general(df)

    if agenda_general.empty:
        data = pd.DataFrame(
            {"mensaje": ["No se encontraron arpónes críticos ni alarmados", "Todas las señales están en estado NORMAL"]}
        )
        num_casos = 0
    else:
        data = agenda_general
        num_casos = agenda_general.shape[0]

    return {
        "data": data,
        "num_casos": num_casos,
        "excel": generar_excel_atenuaciones(df),
    }
