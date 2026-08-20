import io
import math
from numbers import Real

import pandas as pd
from openpyxl import Workbook, load_workbook

from funciones import (
    es_error_fill_openpyxl,
    obtener_bytes_archivo,
    reparar_fills_estilos_xlsx,
)


COL_ABONADO = "N\u00b0 Abonado"
COL_EQUIPO_MAC = "EQUIPO MAC"
COL_NSN = "NSN"
ABONADO_ALIASES = [
    "N? Abonado",
    "N Abonado",
    "No Abonado",
    "Nro Abonado",
    "N\u00ba Abonado",
]
PREVIEW_LIMIT = 200
CSV_CHUNK_SIZE = 25000

COLUMNAS_RESULTADO = [
    "n\u00b0 abonado",
    "documento_x",
    "nombre_x",
    "apellido_x",
    "estatus",
    "observaciones",
    "sn",
    "olt",
    "catv",
    "administrative status",
    "status",
]


def _normalizar_nombre_columna(columna):
    if columna is None:
        return ""

    texto = str(columna)
    texto = texto.replace("\ufeff", "").replace("\u00c2", "").replace("\u00ba", "\u00b0")
    return " ".join(texto.strip().casefold().split())


def _nombres_normalizados(nombres):
    return {_normalizar_nombre_columna(nombre) for nombre in nombres}


def _resolver_columnas(columnas, esquema, nombre_archivo):
    disponibles = {}
    for columna in columnas:
        nombre = _normalizar_nombre_columna(columna)
        if nombre and nombre not in disponibles:
            disponibles[nombre] = columna

    resultado = {}
    faltantes = []
    for destino, aliases in esquema.items():
        candidatos = _nombres_normalizados([destino, *aliases])
        encontrada = next(
            (disponibles[candidato] for candidato in candidatos if candidato in disponibles),
            None,
        )
        if encontrada is None:
            faltantes.append(destino)
        else:
            resultado[destino] = encontrada

    if faltantes:
        raise ValueError(
            "El archivo "
            f"{nombre_archivo} no contiene las columnas requeridas: {', '.join(faltantes)}"
        )

    return resultado


def _resolver_indices_excel(encabezados, esquema, nombre_archivo):
    columnas = _resolver_columnas(encabezados, esquema, nombre_archivo)
    indices_por_columna = {
        columna: indice
        for indice, columna in enumerate(encabezados)
    }
    return {
        destino: indices_por_columna[columna]
        for destino, columna in columnas.items()
    }


def _stream_archivo(archivo):
    stream = getattr(archivo, "stream", archivo)
    if hasattr(stream, "seek"):
        stream.seek(0)
    return stream


def _abrir_excel_streaming(archivo):
    stream = _stream_archivo(archivo)

    try:
        return load_workbook(
            stream,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as error:
        if not es_error_fill_openpyxl(error):
            raise

    reparado = reparar_fills_estilos_xlsx(obtener_bytes_archivo(archivo))
    return load_workbook(
        io.BytesIO(reparado),
        read_only=True,
        data_only=True,
        keep_links=False,
    )


def _iterar_filas_excel(archivo, esquema, nombre_archivo):
    workbook = _abrir_excel_streaming(archivo)
    try:
        worksheet = workbook.active
        filas = worksheet.iter_rows(values_only=True)
        encabezados = next(filas, None)
        if encabezados is None:
            raise ValueError(f"El archivo {nombre_archivo} esta vacio.")

        indices = _resolver_indices_excel(encabezados, esquema, nombre_archivo)
        for fila in filas:
            yield {
                destino: fila[indice] if indice < len(fila) else None
                for destino, indice in indices.items()
            }
    finally:
        workbook.close()


def _es_vacio(valor):
    if valor is None or valor is pd.NA:
        return True
    if isinstance(valor, float) and math.isnan(valor):
        return True
    if isinstance(valor, str) and not valor.strip():
        return True
    return False


def _texto(valor):
    if _es_vacio(valor):
        return ""
    if isinstance(valor, Real) and not isinstance(valor, bool):
        numero = float(valor)
        if numero.is_integer():
            return str(int(numero))
    return str(valor).strip()


def _llave(valor):
    texto = _texto(valor)
    if not texto:
        return None
    if texto.endswith(".0") and texto[:-2].isdigit():
        return texto[:-2]
    return texto


def _ultimos_ocho_caracteres(valor):
    texto = _texto(valor)
    return texto[-8:] if texto else None


def _valor_salida(valor):
    return None if _es_vacio(valor) else valor


def _leer_cortes(archivo):
    esquema = {
        COL_ABONADO: ABONADO_ALIASES,
        COL_EQUIPO_MAC: [],
        "observaciones": [],
    }
    cortes = []
    abonados = set()

    for fila in _iterar_filas_excel(archivo, esquema, "de corte de abonados"):
        if not _es_vacio(fila["observaciones"]):
            continue

        abonado_key = _llave(fila[COL_ABONADO])
        equipo_key = _ultimos_ocho_caracteres(fila[COL_EQUIPO_MAC])
        if abonado_key is None or equipo_key is None:
            continue

        cortes.append(
            {
                "abonado_key": abonado_key,
                "equipo_key": equipo_key,
                "n\u00b0 abonado": _valor_salida(fila[COL_ABONADO]),
                "observaciones": _valor_salida(fila["observaciones"]),
            }
        )
        abonados.add(abonado_key)

    return cortes, abonados


def _leer_saeplus(archivo, abonados_necesarios):
    if not abonados_necesarios:
        return {}

    esquema = {
        COL_ABONADO: ABONADO_ALIASES,
        "documento": [],
        "nombre": [],
        "apellido": [],
        "estatus": ["Estatus", "estatus sae", "estatus saeplus"],
    }
    saeplus = {}

    for fila in _iterar_filas_excel(archivo, esquema, "de abonados SAEPlus"):
        abonado_key = _llave(fila[COL_ABONADO])
        if abonado_key not in abonados_necesarios:
            continue

        estatus = _texto(fila["estatus"])
        if not estatus or estatus.casefold() == "activo":
            continue

        saeplus.setdefault(abonado_key, []).append(
            {
                "documento_x": _valor_salida(fila["documento"]),
                "nombre_x": _valor_salida(fila["nombre"]),
                "apellido_x": _valor_salida(fila["apellido"]),
                "estatus": _valor_salida(fila["estatus"]),
            }
        )

    return saeplus


def _columnas_csv(archivo, esquema, nombre_archivo):
    stream = _stream_archivo(archivo)
    encabezados = pd.read_csv(stream, nrows=0).columns
    columnas = _resolver_columnas(encabezados, esquema, nombre_archivo)
    if hasattr(stream, "seek"):
        stream.seek(0)
    return columnas


def _serie_texto(serie):
    return serie.fillna("").astype(str).str.strip()


def _leer_olt_activos(archivo, equipos_necesarios):
    if not equipos_necesarios:
        return {}

    esquema = {
        "sn": ["SN"],
        "olt": [],
        "catv": [],
        "administrative status": [],
        "status": [],
    }
    columnas = _columnas_csv(archivo, esquema, "de SmartOLT")
    rename_map = {origen: destino for destino, origen in columnas.items()}
    olt = {}

    stream = _stream_archivo(archivo)
    for chunk in pd.read_csv(
        stream,
        dtype=object,
        low_memory=False,
        usecols=list(columnas.values()),
        chunksize=CSV_CHUNK_SIZE,
    ):
        chunk = chunk.rename(columns=rename_map)
        chunk[COL_NSN] = chunk["sn"].map(_ultimos_ocho_caracteres)
        chunk = chunk[chunk[COL_NSN].isin(equipos_necesarios)]
        if chunk.empty:
            continue

        catv = _serie_texto(chunk["catv"]).str.casefold()
        status = _serie_texto(chunk["status"]).str.casefold()
        administrative_status = _serie_texto(chunk["administrative status"]).str.casefold()
        activos = (
            (status == "online")
            | (catv != "disabled")
            | (administrative_status == "enabled")
        )
        chunk = chunk[activos]

        for datos in chunk.to_dict(orient="records"):
            nsn = datos[COL_NSN]
            olt.setdefault(nsn, []).append(
                {
                    "sn": _valor_salida(datos["sn"]),
                    "olt": _valor_salida(datos["olt"]),
                    "catv": _valor_salida(datos["catv"]),
                    "administrative status": _valor_salida(datos["administrative status"]),
                    "status": _valor_salida(datos["status"]),
                }
            )

    return olt


def _iterar_resultados(cortes, saeplus, olt):
    for corte in cortes:
        registros_saeplus = saeplus.get(corte["abonado_key"])
        if not registros_saeplus:
            continue

        registros_olt = olt.get(corte["equipo_key"])
        if not registros_olt:
            continue

        for saeplus_row in registros_saeplus:
            for olt_row in registros_olt:
                yield {
                    "n\u00b0 abonado": corte["n\u00b0 abonado"],
                    "documento_x": saeplus_row["documento_x"],
                    "nombre_x": saeplus_row["nombre_x"],
                    "apellido_x": saeplus_row["apellido_x"],
                    "estatus": saeplus_row["estatus"],
                    "observaciones": corte["observaciones"],
                    "sn": olt_row["sn"],
                    "olt": olt_row["olt"],
                    "catv": olt_row["catv"],
                    "administrative status": olt_row["administrative status"],
                    "status": olt_row["status"],
                }


def _construir_salida(cortes, saeplus, olt):
    output = io.BytesIO()
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Resultado Filtrado")
    worksheet.append(COLUMNAS_RESULTADO)

    preview = []
    total = 0
    for fila in _iterar_resultados(cortes, saeplus, olt):
        total += 1
        worksheet.append([fila[columna] for columna in COLUMNAS_RESULTADO])
        if len(preview) < PREVIEW_LIMIT:
            preview.append(fila)

    workbook.save(output)
    output.seek(0)

    return {
        "data": pd.DataFrame(preview, columns=COLUMNAS_RESULTADO),
        "columns": COLUMNAS_RESULTADO,
        "num_casos": total,
        "excel": output,
    }


def procesar_cortes(abonados_file, cortes_file, sae_file):
    cortes, abonados_necesarios = _leer_cortes(abonados_file)
    saeplus = _leer_saeplus(sae_file, abonados_necesarios)
    equipos_necesarios = {
        corte["equipo_key"]
        for corte in cortes
        if corte["abonado_key"] in saeplus
    }
    olt = _leer_olt_activos(cortes_file, equipos_necesarios)

    return _construir_salida(cortes, saeplus, olt)
